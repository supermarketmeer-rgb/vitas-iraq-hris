# -*- coding: utf-8 -*-
"""
Script to generate an executive, persuasive Sales & Value Proposition PowerPoint (.pptx)
for the HRIS (Human Resources Information System) targeting C-level executives, HR directors, and business owners.
"""

import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation(output_path="HRIS_System_Presentation.pptx"):
    prs = Presentation()
    # 16:9 Widescreen standard
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    blank_slide_layout = prs.slide_layouts[6] # Blank layout
    
    # Color Palette
    COLOR_PRIMARY = RGBColor(15, 23, 42)      # Deep Navy / Slate 900
    COLOR_SECONDARY = RGBColor(13, 148, 136)  # Teal 600
    COLOR_ACCENT = RGBColor(217, 119, 6)     # Amber 600
    COLOR_BG_LIGHT = RGBColor(248, 250, 252) # Slate 50
    COLOR_CARD_BG = RGBColor(255, 255, 255)  # White
    COLOR_CARD_BORDER = RGBColor(226, 232, 240) # Slate 200
    COLOR_TEXT_DARK = RGBColor(30, 41, 59)   # Slate 800
    COLOR_TEXT_MUTED = RGBColor(100, 116, 139) # Slate 500
    COLOR_SUCCESS = RGBColor(16, 185, 129)   # Emerald 500

    def add_header(slide, title_text, category_text="نظام إدارة الموارد البشرية الشامل (HRIS)"):
        # Header background banner
        header_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.2))
        header_box.fill.solid()
        header_box.fill.fore_color.rgb = COLOR_PRIMARY
        header_box.line.color.rgb = COLOR_PRIMARY

        # Accent bottom line
        accent_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.15), Inches(13.333), Inches(0.06))
        accent_line.fill.solid()
        accent_line.fill.fore_color.rgb = COLOR_SECONDARY
        accent_line.line.color.rgb = COLOR_SECONDARY

        # Category text
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(11.7), Inches(0.3))
        cat_tf = cat_box.text_frame
        cat_tf.word_wrap = True
        p_cat = cat_tf.paragraphs[0]
        p_cat.text = category_text
        p_cat.alignment = PP_ALIGN.RIGHT
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_ACCENT

        # Title text
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.42), Inches(11.7), Inches(0.65))
        title_tf = title_box.text_frame
        title_tf.word_wrap = True
        p_title = title_tf.paragraphs[0]
        p_title.text = title_text
        p_title.alignment = PP_ALIGN.RIGHT
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = RGBColor(255, 255, 255)

    def set_slide_background(slide, color=COLOR_BG_LIGHT):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()

    def add_card(slide, left, top, width, height, title, body_bullets, badge=None, badge_color=COLOR_SECONDARY):
        # Card Background
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD_BG
        card.line.color.rgb = COLOR_CARD_BORDER
        card.line.width = Pt(1)

        # Header Accent
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left + width - Inches(0.12), top + Inches(0.2), Inches(0.08), height - Inches(0.4))
        accent.fill.solid()
        accent.fill.fore_color.rgb = badge_color
        accent.line.color.rgb = badge_color

        tb = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), width - Inches(0.55), height - Inches(0.35))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        # Badge if any
        if badge:
            p_badge = tf.paragraphs[0]
            p_badge.text = f"[{badge}]"
            p_badge.font.size = Pt(11)
            p_badge.font.bold = True
            p_badge.font.color.rgb = badge_color
            p_badge.alignment = PP_ALIGN.RIGHT
            p_title = tf.add_paragraph()
        else:
            p_title = tf.paragraphs[0]

        p_title.text = title
        p_title.font.size = Pt(16)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_TEXT_DARK
        p_title.alignment = PP_ALIGN.RIGHT
        p_title.space_after = Pt(8)

        for bullet in body_bullets:
            p_b = tf.add_paragraph()
            p_b.text = f"•  {bullet}"
            p_b.font.size = Pt(12)
            p_b.font.color.rgb = COLOR_TEXT_MUTED
            p_b.alignment = PP_ALIGN.RIGHT
            p_b.space_after = Pt(4)

    # ==========================================
    # SLIDE 1: Cover Slide (غلاف العرض التنفيذي)
    # ==========================================
    s1 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s1, COLOR_PRIMARY)

    # Decorative shapes
    dec1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.4), Inches(7.5))
    dec1.fill.solid()
    dec1.fill.fore_color.rgb = COLOR_SECONDARY
    dec1.line.fill.background()

    dec2 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(0), Inches(0.15), Inches(7.5))
    dec2.fill.solid()
    dec2.fill.fore_color.rgb = COLOR_ACCENT
    dec2.line.fill.background()

    # Cover Text Box
    cov_tb = s1.shapes.add_textbox(Inches(1.5), Inches(1.8), Inches(10.5), Inches(4.5))
    cov_tf = cov_tb.text_frame
    cov_tf.word_wrap = True

    p = cov_tf.paragraphs[0]
    p.text = "نظام إدارة الموارد البشرية والرواتب المتكامل (HRIS Enterprise)"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.RIGHT
    p.space_after = Pt(14)

    p2 = cov_tf.add_paragraph()
    p2.text = "الحل التقني الاستراتيجي للتحول الرقمي، ضبط التكاليف، والامتثال القانوني الكامل"
    p2.font.size = Pt(18)
    p2.font.color.rgb = COLOR_SECONDARY
    p2.alignment = PP_ALIGN.RIGHT
    p2.space_after = Pt(28)

    p3 = cov_tf.add_paragraph()
    p3.text = "عرض القيمة الاستثمارية وكفاءة التشغيل المقدم إلى الإدارة العليا وصناع القرار"
    p3.font.size = Pt(14)
    p3.font.color.rgb = RGBColor(203, 213, 225) # Slate 300
    p3.alignment = PP_ALIGN.RIGHT
    p3.space_after = Pt(20)

    p4 = cov_tf.add_paragraph()
    p4.text = "✓ محرك عقود ذكي وديناميكي   |   ✓ رواتب وبصمة مؤتمتة 100%   |   ✓ أمان وسرعة محلية مع مزامنة سحابية"
    p4.font.size = Pt(13)
    p4.font.bold = True
    p4.font.color.rgb = COLOR_ACCENT
    p4.alignment = PP_ALIGN.RIGHT

    # ==========================================
    # SLIDE 2: التحديات الإدارية والمخاطر التشغيلية (Pain Points)
    # ==========================================
    s2 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s2)
    add_header(s2, "واقع إدارة الموارد البشرية التقليدية: التحديات والتكاليف الخفية")

    col_w = Inches(3.65)
    gap = Inches(0.35)
    top_pos = Inches(1.6)
    h_pos = Inches(5.3)

    add_card(s2, Inches(9.1), top_pos, col_w, h_pos, 
             "1. نزيف الوقت والأخطاء اليدوية",
             [
                 "إعداد الرواتب في ملفات Excel يستغرق أياماً ويتعرض لأخطاء حسابية مكلفة.",
                 "صياغة العقود وتعديل بنودها يدوياً لكل موظف يهدر عشرات الساعات.",
                 "ضياع أرشيف الوثائق وصعوبة الوصول الفوري لملف الموظف عند الطلب.",
                 "تسجيل الإجازات ورقياً وغياب الرؤية المباشرة للأرصدة المستحقة."
             ], badge="كفاءة متدنية", badge_color=RGBColor(239, 68, 68))

    add_card(s2, Inches(5.1), top_pos, col_w, h_pos,
             "2. مخاطر الامتثال وغرامات العمل",
             [
                 "صعوبة احتساب الضمان الاجتماعي وضرائب الدخل وفق القوانين الرسمية.",
                 "عقود عمل غير منضبطة قانونياً تعرض المؤسسة لقضايا ونزاعات عمالية.",
                 "غياب سجل تدقيق (Audit Trail) لمعرفة من قام بتعديل راتب أو رصيد.",
                 "عشوائية تطبيق الجزاءات والخصومات ومخالفتها لتعليمات وزارة العمل."
             ], badge="مخاطر قانونية", badge_color=RGBColor(245, 158, 11))

    add_card(s2, Inches(1.1), top_pos, col_w, h_pos,
             "3. انقطاع الإنترنت وضعف الحماية",
             [
                 "الأنظمة السحابية البحتة تتعطل تماماً فور انقطاع الإنترنت في الفروع.",
                 "تخوف الإدارة من وضع بيانات الرواتب الحساسة في خوادم مجهولة أو بطيئة.",
                 "صعوبة دمج أجهزة البصمة الميدانية مع قواعد البيانات المركزية.",
                 "غياب تحليلات فورية لصانع القرار حول تكلفة القوى العاملة ومعدل الدوران."
             ], badge="عائق تقني", badge_color=RGBColor(99, 102, 241))

    # ==========================================
    # SLIDE 3: الحل المتكامل (Integrated Solution Overview)
    # ==========================================
    s3 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s3)
    add_header(s3, "الحل المتكامل: منظومة HRIS الشاملة من الألف إلى الياء")

    w4 = Inches(2.7)
    g4 = Inches(0.25)
    t4 = Inches(1.6)
    h4 = Inches(5.3)

    add_card(s3, Inches(9.9), t4, w4, h4,
             "ملف الموظف وعقود العمل",
             [
                 "سجل رقمي متكامل 360 درجة لكل موظف.",
                 "محرك نماذج العقود وبنودها القابلة للتخصيص.",
                 "استبدال تلقائي للبيانات والرواتب في نصوص العقود.",
                 "أرشفة المستندات والبطاقات الوطنية بضغطة زر."
             ], badge="إدارة الأفراد", badge_color=COLOR_SECONDARY)

    add_card(s3, Inches(7.0), t4, w4, h4,
             "حضور ذكي وبصمة حية",
             [
                 "ربط مباشر مع أجهزة البصمة (ZKTeco/TCP).",
                 "احتساب التأخير، المغادرة المبكرة، والإضافي آلياً.",
                 "مركز موافقات إلكتروني لطلبات الإجازة والغياب.",
                 "رصيد إجازات ذكي محدث في الوقت الفعلي."
             ], badge="الانضباط والوقت", badge_color=COLOR_SECONDARY)

    add_card(s3, Inches(4.1), t4, w4, h4,
             "محرك الرواتب والضرائب",
             [
                 "توليد كشوفات الرواتب لآلاف الموظفين بدقائق.",
                 "احتساب الضمان الاجتماعي والاستقطاعات الضريبية.",
                 "كشوفات بنكية جاهزة للتحويل المالي المباشر.",
                 "قسائم رواتب سرية وإشعارات للموظفين."
             ], badge="المالية والحسابات", badge_color=COLOR_SECONDARY)

    add_card(s3, Inches(1.2), t4, w4, h4,
             "المعمارية الهجينة والأمان",
             [
                 "سرعة تشغيل محلية خارقة مع دعم المزامنة السحابية.",
                 "العمل يستمر 100% حتى مع انقطاع الإنترنت الكامل.",
                 "صلاحيات أمنية دقيقة بحسب القسم والمستوى الوظيفي.",
                 "سجل تتبع ومراقبة العمليات (Audit Logging)."
             ], badge="الاستقرار التقني", badge_color=COLOR_ACCENT)

    # ==========================================
    # SLIDE 4: الميزة التنافسية - العقود الذكية والطباعة
    # ==========================================
    s4 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s4)
    add_header(s4, "محرك إدارة العقود الذكي: توثيق قانوني وتوليد فوري بضغطة زر")

    add_card(s4, Inches(7.0), Inches(1.6), Inches(5.5), Inches(5.3),
             "كيف يحول النظام إدارة العقود إلى عملية مؤتمتة بالكامل؟",
             [
                 "مكتبة نماذج مرنة (Contract Templates): إنشاء وتعديل نماذج عقود متعددة (عقد محدد المدة، غير محدد، تجريبي، استشاري) مع صلاحية كاملة لتعديل البنود والشروط في أي وقت من شاشة الإعدادات.",
                 "محرك المتغيرات الديناميكية (Smart Placeholders): يقوم النظام باستبدال المتغيرات آلياً بنصوص حقيقية مثل اسم الموظف، المسمى الوظيفي، الراتب الأساسي، تاريخ المباشرة، دون أي تدخل يدوي.",
                 "معاينة وطباعة A4 فورية: نافذة عرض متطورة تطبق معايير الطباعة الرسمية للشركات، مع أزرار مباشرة لتحميل العقد كـ PDF أو إرساله للطابعة.",
                 "تحديث فوري للبنود (Real-time Clauses Sync): عند تعديل أي بند في الإعدادات، ينعكس فوراً وبشكل دقيق على عقود الموظفين دون تكرار أو تشويه للتنسيق."
             ], badge="محرك العقود", badge_color=COLOR_SECONDARY)

    add_card(s4, Inches(0.8), Inches(1.6), Inches(5.8), Inches(5.3),
             "العائد المباشر على المنشأة من ميزة العقود",
             [
                 "توفير 95% من وقت قسم الموارد البشرية المخصص لصياغة وتدقيق العقود الفردية.",
                 "القضاء التام على أخطاء الرواتب أو المسميات المكتوبة في العقود الورقية القديمة.",
                 "توحيد الصيغة القانونية المعتمدة للشركة وحماية حقوق العمل وواجبات الموظف.",
                 "أرشفة رقمية آمنة تتيح استدعاء عقد أي موظف خلال ثانيتين فقط عند التفتيش أو التدقيق الداخلي.",
                 "جاهزية كاملة لطباعة عقود الدفعات الكبيرة من التعيينات الجديدة دفعة واحدة."
             ], badge="قيمة ملموسة للعميل", badge_color=COLOR_ACCENT)

    # ==========================================
    # SLIDE 5: نظام الحضور والانصراف والربط البيومتري
    # ==========================================
    s5 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s5)
    add_header(s5, "الحضور والانصراف الذكي: ربط أجهزة البصمة وأتمتة الانضباط")

    add_card(s5, Inches(9.1), Inches(1.6), col_w, Inches(5.3),
             "سحب تلقائي للبصمات",
             [
                 "دعم بروتوكولات الاتصال المباشر مع أجهزة البصمة وبصمة الوجه (ZKTeco وغيرها).",
                 "تفريغ حركات الدخول والخروج في الوقت الحقيقي دون الاعتماد على الفلاش ميموري.",
                 "كشف تلقائي لحالات البصمة المزدوجة أو النسيان وتصحيحها ذكياً."
             ], badge="ربط عتادي مباشر", badge_color=COLOR_SECONDARY)

    add_card(s5, Inches(5.1), Inches(1.6), col_w, Inches(5.3),
             "مركز الموافقات الإلكتروني",
             [
                 "شاشة موحدة لمدراء الأقسام لمراجعة واعتماد الإجازات والمغادرات.",
                 "مقارنة فورية بين رصيد الإجازة المتبقي والطلب المقدم قبل الموافقة.",
                 "إشعارات واضحة تحدد مسار الموافقة وتسلسل الصلاحيات الإدارية."
             ], badge="سير عمل لا ورقي", badge_color=COLOR_SECONDARY)

    add_card(s5, Inches(1.1), Inches(1.6), col_w, Inches(5.3),
             "ربط مباشر بالرواتب",
             [
                 "ترحيل تلقائي لخصومات التأخير وساعات العمل الإضافي إلى مسيرات الرواتب.",
                 "القضاء على ظاهرة التحيز أو المحسوبية في تطبيق قواعد الخصم والمكافأة.",
                 "تقارير دوام تفصيلية وملخصة تساعد الإدارة على تقييم التزام الكادر."
             ], badge="عدالة ودقة حسابية", badge_color=COLOR_ACCENT)

    # ==========================================
    # SLIDE 6: محرك الرواتب والامتثال القانوني والضريبي
    # ==========================================
    s6 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s6)
    add_header(s6, "محرك الرواتب والامتثال: صفر أخطاء وامتثال كامل للتشريعات")

    add_card(s6, Inches(7.0), Inches(1.6), Inches(5.5), Inches(5.3),
             "أتمتة الحسابات المالية المعقدة",
             [
                 "احتساب الراتب الإجمالي والصافي، والبدلات (سكن، نقل، طعام، مهنية) بشكل آلي.",
                 "معادلات ديناميكية لاحتساب ضريبة الدخل السنوية والشهرية وفق الشرائح الرسمية.",
                 "تطبيق استقطاعات الضمان الاجتماعي (حصة الموظف وحصة الشركة) بدقة متناهية.",
                 "إدارة السلف والقروض وجدولتها على أقساط تخصم تلقائياً من المسير الشهري.",
                 "إصدار كشف البنك الموحد وتصديره بصيغة Excel المتوافقة مع بوابات التحويل البنكي."
             ], badge="دقة مالية 100%", badge_color=COLOR_SECONDARY)

    add_card(s6, Inches(0.8), Inches(1.6), Inches(5.8), Inches(5.3),
             "ماذا يكسب المدير المالي (CFO) وإدارة الشركة؟",
             [
                 "تقليص مدة إغلاق الرواتب الشهرية من 4 أيام عمل إلى أقل من 30 دقيقة فقط.",
                 "تجنب الغرامات المالية والعقوبات القانونية الناتجة عن أخطاء الضمان والضريبة.",
                 "شفافية مطلقة أمام الموظفين عبر قسيمة الراتب الرقمية (Digital Payslip).",
                 "حماية أمنية مشددة وسرية مطلقة لبيانات الرواتب عبر صلاحيات وصول مشفرة ومقيدة.",
                 "مقارنة تكلفة الرواتب عبر الشهور وتحليل الميزانيات التقديرية بدقة."
             ], badge="أثر مالي مباشر", badge_color=COLOR_SUCCESS)

    # ==========================================
    # SLIDE 7: المعمارية الهجينة (Hybrid Cloud + Local)
    # ==========================================
    s7 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s7)
    add_header(s7, "المعمارية الهجينة الذكية: استمرارية عمل لا تتوقف مع سرعة وأمان")

    add_card(s7, Inches(7.0), Inches(1.6), Inches(5.5), Inches(5.3),
             "المشكلة الكبرى للأنظمة السحابية والحل المبتكر في نظامنا",
             [
                 "التحدي في السوق المحلي: انقطاع شبكة الإنترنت أو بطء خطوط الاتصال يؤدي لشلل كامل في عمل الإدارات والمصانع والفروع إذا كان النظام سحابياً فقط.",
                 "معمارية نظامنا المبتكرة (Hybrid Engine):",
                 "  1. أداء محلي فائق السرعة (Local First): يعمل النظام على الشبكة الداخلية بأقصى سرعة استجابة (Zero Latency) حتى لو انقطعت خدمة الإنترنت لأيام.",
                 "  2. مزامنة سحابية تلقائية ذكية (Auto Cloud Sync): فور توفر الاتصال، يقوم النظام بمزامنة البيانات المشفرة إلى السحابة المركزية أو قاعدة بيانات الفروع بأمان كامل.",
                 "  3. نسخ احتياطي متعدد (Fail-safe Backup): لا مجال لضياع أي سجل أو حركة حضور أو بيانات مالية مهما كانت الظروف."
             ], badge="حصرية هندسية", badge_color=COLOR_SECONDARY)

    add_card(s7, Inches(0.8), Inches(1.6), Inches(5.8), Inches(5.3),
             "لماذا يفضل مسؤولو تقنية المعلومات (IT Directors) نظامنا؟",
             [
                 "تحكم كامل في البيانات (Data Sovereignty): إمكانية استضافة البيانات محلياً داخل خوادم الشركة أو على سحابة خاصة.",
                 "استهلاك منخفض لموارد الخادم وخفة تشغيل عالية بالاعتماد على أحدث التقنيات الحديثة (Node.js, SQLite/PostgreSQL, React).",
                 "سهولة الصيانة والتحديث المركزي دون تعطيل عمل موظفي الموارد البشرية.",
                 "توافق مع مختلف بيئات العمل: أجهزة الكمبيوتر، الأجهزة اللوحية، والهواتف الذكية."
             ], badge="اعتمادية تقنية فائقة", badge_color=COLOR_ACCENT)

    # ==========================================
    # SLIDE 8: العائد على الاستثمار والوفر المالي الملموس (ROI)
    # ==========================================
    s8 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s8)
    add_header(s8, "العائد المالي على الاستثمار (ROI): النظام يدفع تكلفته بنفسه")

    # 4 KPI Cards
    kpi_w = Inches(2.7)
    kpi_h = Inches(2.3)
    
    def add_kpi(slide, left, top, value, label, subtext, color=COLOR_SECONDARY):
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, kpi_w, kpi_h)
        box.fill.solid()
        box.fill.fore_color.rgb = COLOR_CARD_BG
        box.line.color.rgb = color
        box.line.width = Pt(2)

        tf = box.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = value
        p.font.size = Pt(30)
        p.font.bold = True
        p.font.color.rgb = color
        p.alignment = PP_ALIGN.CENTER
        
        p2 = tf.add_paragraph()
        p2.text = label
        p2.font.size = Pt(14)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_TEXT_DARK
        p2.alignment = PP_ALIGN.CENTER
        p2.space_after = Pt(4)

        p3 = tf.add_paragraph()
        p3.text = subtext
        p3.font.size = Pt(11)
        p3.font.color.rgb = COLOR_TEXT_MUTED
        p3.alignment = PP_ALIGN.CENTER

    add_kpi(s8, Inches(9.9), Inches(1.6), "75%+", "توفير وقت قسم الـ HR", "أتمتة الرواتب، العقود، والإجازات")
    add_kpi(s8, Inches(7.0), Inches(1.6), "0%", "نسبة أخطاء الحسابات", "محرك رواتب ومطابقة حضور آلي")
    add_kpi(s8, Inches(4.1), Inches(1.6), "100%", "التخلص من الهدر الورقي", "أرشفة رقمية وعقود إلكترونية موثقة")
    add_kpi(s8, Inches(1.2), Inches(1.6), "< 3 أشهر", "فترة استرداد الاستثمار", "وفر مباشر في التكاليف التشغيلية")

    # Detailed Summary Card below KPIs
    roi_card = slide = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(4.2), Inches(11.4), Inches(2.7))
    roi_card.fill.solid()
    roi_card.fill.fore_color.rgb = COLOR_PRIMARY
    roi_card.line.color.rgb = COLOR_SECONDARY
    roi_card.line.width = Pt(1.5)

    rtf = roi_card.text_frame
    rtf.word_wrap = True
    
    rp1 = rtf.paragraphs[0]
    rp1.text = "معادلة الوفر المالي التقديري لشركة تضم 100 موظف:"
    rp1.font.size = Pt(17)
    rp1.font.bold = True
    rp1.font.color.rgb = COLOR_ACCENT
    rp1.alignment = PP_ALIGN.RIGHT
    rp1.space_after = Pt(8)

    rp2 = rtf.add_paragraph()
    rp2.text = "• وفر ساعات العمل الإدارية: توفير ما لا يقل عن 60 ساعة شهرياً كانت تضيع في تجميع بصمات الدوام وصياغة العقود وتعديل جداول الرواتب (وفر مالي يتجاوز 1,500$ شهرياً)."
    rp2.font.size = Pt(13)
    rp2.font.color.rgb = RGBColor(241, 245, 249)
    rp2.alignment = PP_ALIGN.RIGHT
    rp2.space_after = Pt(6)

    rp3 = rtf.add_paragraph()
    rp3.text = "• منع التسرب المالي: ضبط دقيق للدقائق الضائعة، الغيابات غير الموثقة، وأخطاء حساب ساعات الإضافي غير المصرح بها."
    rp3.font.size = Pt(13)
    rp3.font.color.rgb = RGBColor(241, 245, 249)
    rp3.alignment = PP_ALIGN.RIGHT
    rp3.space_after = Pt(6)

    rp4 = rtf.add_paragraph()
    rp4.text = "• حماية السمعة والتكاليف القضائية: تجنب النزاعات العمالية والغرامات المفروضة من لجان تفتيش العمل بسبب العقود غير المطابقة أو تأخر صرف المستحقات."
    rp4.font.size = Pt(13)
    rp4.font.color.rgb = RGBColor(241, 245, 249)
    rp4.alignment = PP_ALIGN.RIGHT

    # ==========================================
    # SLIDE 9: تجربة المستخدم وسهولة الاستخدام (UI/UX)
    # ==========================================
    s9 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s9)
    add_header(s9, "واجهة استخدام عصرية: تبسيط العمليات وتجربة عمل مريحة")

    add_card(s9, Inches(7.0), Inches(1.6), Inches(5.5), Inches(5.3),
             "تصميم يلغي الحاجة لتدريب طويل ومعقد",
             [
                 "واجهة عربية كاملة مصممة وفق أحدث معايير تجربة المستخدم (Enterprise UX).",
                 "لوحات مؤشرات قيادية (Dashboards) تفاعلية تعرض حالة القوى العاملة، الغيابات، وتوزيع الرواتب بلحظة واحدة.",
                 "نظام فلترة وبحث فوري وسريع بين آلاف الموظفين والسجلات.",
                 "إمكانية العمل بالوضع المظلم (Dark Mode) أو الفاتح حسب تفضيل المستخدم ومستوى الإضاءة.",
                 "تدرج بصري وألوان متناسقة تمنح إحساساً بالفخامة والاحترافية العالية."
             ], badge="بساطة وسرعة", badge_color=COLOR_SECONDARY)

    add_card(s9, Inches(0.8), Inches(1.6), Inches(5.8), Inches(5.3),
             "ماذا يعني ذلك لمدير الموارد البشرية والموظفين؟",
             [
                 "اعتماد فوري للنظام من قبل موظفي الموارد البشرية خلال أيام معدودة دون مقاومة للتغيير.",
                 "تقليص أخطاء الإدخال بفضل الحقول الذكية والتنبيهات المسبقة عند وجود تعارض في البيانات.",
                 "تصدير واستيراد مرن لكافة التقارير بصيغ Excel وPDF جاهزة للعرض الفوري في اجتماعات مجلس الإدارة.",
                 "مرونة في الوصول من مختلف الأجهزة داخل الشبكة دون الحاجة لتثبيت برمجيات معقدة على كل جهاز."
             ], badge="تبني سلس للنظام", badge_color=COLOR_ACCENT)

    # ==========================================
    # SLIDE 10: خطة التطبيق السريع والدعم الفني (Implementation Roadmap)
    # ==========================================
    s10 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s10)
    add_header(s10, "خطة التحول الرقمي: تطبيق سلس، تدريب شامل، ودعم مستمر")

    step_w = Inches(2.7)
    step_h = Inches(5.3)

    add_card(s10, Inches(9.9), Inches(1.6), step_w, step_h,
             "المرحلة 1: التهيئة والبيانات",
             [
                 "استيراد بيانات الموظفين الحالية من ملفات Excel السابقة.",
                 "تهيئة الهيكل التنظيمي، الأقسام، والمسميات الوظيفية.",
                 "إعداد نماذج العقود المعتمدة وشروطها الخاصة."
             ], badge="الأسبوع 1", badge_color=COLOR_SECONDARY)

    add_card(s10, Inches(7.0), Inches(1.6), step_w, step_h,
             "المرحلة 2: ربط البصمة",
             [
                 "ربط أجهزة البصمة بالفروع والتأكد من استقرار تدفق البيانات.",
                 "ضبط جداول الورديات، أوقات السماح، وسياسات التأخير.",
                 "إجراء تجارب محاكاة لعملية البصمة اليومية."
             ], badge="الأسبوع 2", badge_color=COLOR_SECONDARY)

    add_card(s10, Inches(4.1), Inches(1.6), step_w, step_h,
             "المرحلة 3: التدريب والتشغيل",
             [
                 "تدريب عملي مكثف لفريق الموارد البشرية والإدارة المالية.",
                 "إصدار أول مسير رواتب تجريبي ومطابقته بدقة 100%.",
                 "تفعيل صلاحيات المستخدمين والمدراء المباشرين."
             ], badge="الأسبوع 3", badge_color=COLOR_SECONDARY)

    add_card(s10, Inches(1.2), Inches(1.6), step_w, step_h,
             "المرحلة 4: الانطلاق والدعم",
             [
                 "الانطلاق الرسمي للنظام والاستغناء عن المعاملات الورقية.",
                 "دعم فني استباقي ومتابعة يومية لأول شهر تشغيلي.",
                 "تحديثات وتطويرات دورية لضمان استمرارية الريادة."
             ], badge="الأسبوع 4", badge_color=COLOR_SUCCESS)

    # ==========================================
    # SLIDE 11: مقارنة تنافسية (Competitive Advantage)
    # ==========================================
    s11 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s11)
    add_header(s11, "لماذا يختار عملاؤنا نظامنا مقارنة بالبدائل الأخرى في السوق؟")

    add_card(s11, Inches(7.0), Inches(1.6), Inches(5.5), Inches(5.3),
             "مقارنة بنظام ملفات Excel والحلول اليدوية",
             [
                 "الحلول اليدوية: معرضة للتلف، الضياع، التلاعب، ولا توفر أي سرية لبيانات الرواتب، وتستنزف مئات الساعات شهرياً.",
                 "نظامنا الرقمي: قاعدة بيانات مركزية مشفرة، سجل تدقيق لكل نقرة، وحسابات رواتب وبصمات فورية ودقيقة 100% دون أي فرصة للتلاعب البشري."
             ], badge="تفوق كامل على اليدوي", badge_color=COLOR_ACCENT)

    add_card(s11, Inches(0.8), Inches(1.6), Inches(5.8), Inches(5.3),
             "مقارنة بالأنظمة السحابية المستوردة (Global SaaS)",
             [
                 "الأنظمة الأجنبية المستوردة: تكاليف اشتراك شهرية باهظة بالدولار لكل موظف، لا تفهم خصوصية قانون العمل والضرائب المحلية، وتتوقف فور انقطاع الإنترنت.",
                 "نظامنا المحلي المتخصص: صُمم خصيصاً ليطابق القوانين والضرائب المحلية بدقة، يعمل محلياً وسحابياً معاً، مع دعم فني عربي مباشر وتكلفة استثمارية عادلة ومجدية."
             ], badge="تفوق على الأنظمة الأجنبية", badge_color=COLOR_SUCCESS)

    # ==========================================
    # SLIDE 12: الخطوة التالية والدعوة للاتفاق (Closing & Call to Action)
    # ==========================================
    s12 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s12, COLOR_PRIMARY)

    cta_tb = s12.shapes.add_textbox(Inches(1.5), Inches(1.5), Inches(10.5), Inches(5.0))
    cta_tf = cta_tb.text_frame
    cta_tf.word_wrap = True

    cp1 = cta_tf.paragraphs[0]
    cp1.text = "جاهزون للارتقاء بإدارة مواردكم البشرية إلى العصر الرقمي؟"
    cp1.font.size = Pt(28)
    cp1.font.bold = True
    cp1.font.color.rgb = RGBColor(255, 255, 255)
    cp1.alignment = PP_ALIGN.RIGHT
    cp1.space_after = Pt(16)

    cp2 = cta_tf.add_paragraph()
    cp2.text = "ندعوكم لتجربة عرض حي مباشر (Live Interactive Demo) للنظام بكافة مزاياه"
    cp2.font.size = Pt(20)
    cp2.font.bold = True
    cp2.font.color.rgb = COLOR_SECONDARY
    cp2.alignment = PP_ALIGN.RIGHT
    cp2.space_after = Pt(24)

    cp3 = cta_tf.add_paragraph()
    cp3.text = "ما سنقدمه لكم في جلسة العرض التجريبي:"
    cp3.font.size = Pt(16)
    cp3.font.bold = True
    cp3.font.color.rgb = COLOR_ACCENT
    cp3.alignment = PP_ALIGN.RIGHT
    cp3.space_after = Pt(10)

    points = [
        "1. استعراض عملي فوري لتوليد عقد عمل مخصص وطباعته بضغطة زر واحدة.",
        "2. ربط تجريبي لجهاز بصمة وسحب الحركات ومطابقتها على جدول الورديات.",
        "3. احتساب مسير رواتب كامل شامل الاستقطاعات الضريبية والضمان وإصدار قسائم الرواتب.",
        "4. الإجابة عن كافة التساؤلات الفنية وتقديم خطة زمنية مفصلة للتنفيذ تلائم حجم مؤسستكم."
    ]
    for pt in points:
        cp_sub = cta_tf.add_paragraph()
        cp_sub.text = pt
        cp_sub.font.size = Pt(14)
        cp_sub.font.color.rgb = RGBColor(226, 232, 240)
        cp_sub.alignment = PP_ALIGN.RIGHT
        cp_sub.space_after = Pt(6)

    prs.save(output_path)
    print(f"[SUCCESS] Presentation generated successfully at: {output_path}")

if __name__ == "__main__":
    out = os.path.abspath("HRIS_System_Executive_Pitch.pptx")
    create_presentation(out)
